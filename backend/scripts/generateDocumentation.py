"""
Script to generate comprehensive Word (.docx) and PDF (.pdf) documentation
for the InfraPilot Enterprise Infrastructure Management Platform.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

DOCS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../docs'))
DOCX_PATH = os.path.join(DOCS_DIR, 'InfraPilot_Project_Documentation.docx')
PDF_PATH = os.path.join(DOCS_DIR, 'InfraPilot_Project_Documentation.pdf')

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')

def generate_docx():
    print("[DOCX] Generating Microsoft Word (.docx) documentation...")
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Palette
    NAVY = RGBColor(15, 23, 42)      # #0F172A
    BLUE = RGBColor(37, 99, 235)     # #2563EB
    SLATE = RGBColor(71, 85, 105)    # #475569
    DARK_GRAY = RGBColor(30, 41, 59) # #1E293B

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = DARK_GRAY

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_title.add_run("ENTERPRISE ARCHITECTURE & USER MANUAL\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = BLUE
    
    run_title = p_title.add_run("InfraPilot Management Platform")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_desc = p_sub.add_run("Self-Service Infrastructure Provisioning, Governance, & Lifecycle Automation\n")
    run_desc.font.size = Pt(13)
    run_desc.font.italic = True
    run_desc.font.color.rgb = SLATE

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Author: Senior IT Application Developer Candidate  |  Version: 1.0.0 (Production Release)")
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = NAVY
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = BLUE
        return h

    def add_callout(text, title="KEY TAKEAWAY"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_shading(cell, "EFF6FF") # Light blue
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        r1 = p.add_run(f"📌 {title}: ")
        r1.font.bold = True
        r1.font.color.rgb = BLUE
        r2 = p.add_run(text)
        r2.font.color.rgb = DARK_GRAY
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 1. Executive Summary & Purpose
    add_h1("1. Executive Summary & Purpose of the Application")
    doc.add_paragraph(
        "InfraPilot is a production-grade, self-service enterprise infrastructure portal designed to eliminate "
        "the classic tension between developer agility and IT operational governance. In traditional enterprise environments, "
        "software engineering teams often endure multi-day delays waiting for infrastructure teams to manually approve "
        "and provision virtual machines (VMs) through ticketing systems. Conversely, granting developers unfettered direct "
        "access to cloud provider consoles creates severe security risks, untracked configuration drift, and runaway costs."
    )
    doc.add_paragraph(
        "InfraPilot bridges this gap by providing an intuitive, web-based self-service catalog where developers can "
        "request compute infrastructure with automated quota checks. Behind the scenes, requests flow through an enterprise "
        "approval workflow, trigger asynchronous background job queues, execute HashiCorp Terraform automation scripts in "
        "isolated sandboxes, and catalog provisioned virtual machines in an immutable inventory with real-time power controls."
    )
    add_callout(
        "InfraPilot empowers engineers to provision infrastructure on-demand while guaranteeing IT leaders 100% policy compliance, "
        "strict budget quota enforcement, zero double-approvals, and complete audit logging.",
        title="CORE VALUE PROPOSITION"
    )

    # 2. Key Capabilities
    add_h1("2. Core Capabilities & Architectural Highlights")
    doc.add_paragraph("The application delivers six primary pillars of enterprise functionality:")
    
    features = [
        ("Self-Service Catalog", "Standardized VM templates (OS distributions, CPU, RAM, Disk, Region, Environment) validated via Zod schemas."),
        ("Resource Quota Enforcement", "Real-time server-side validation preventing developers from exceeding allocated CPU, RAM, disk, or VM count limits."),
        ("Strict Separation of Duties (RBAC)", "Distinct roles: DEVELOPER (request, track, manage personal VMs) and ADMIN (fleet oversight, quota tuning, audit logs, approvals). Developers cannot approve requests; admins cannot approve their own requests."),
        ("Asynchronous Job Processing", "Dedicated BullMQ engine powered by Redis. HTTP REST requests return immediately while long-running Terraform automation executes reliably in background workers."),
        ("Terraform IaC Sandbox", "Executes terraform init, validate, plan, and apply in isolated directories using the safe hashicorp/local provider to generate JSON infrastructure manifests without cloud charges."),
        ("Fleet Inventory & Audit Trail", "Real-time state tracking with interactive Start/Stop power controls and an append-only compliance audit log recording all logins, approvals, and configuration changes.")
    ]
    for name, desc in features:
        p = doc.add_paragraph(style='List Bullet')
        r1 = p.add_run(f"{name}: ")
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        p.add_run(desc)

    # 3. System Architecture
    add_h1("3. System Architecture & How It Works")
    doc.add_paragraph(
        "InfraPilot adopts a modern Decoupled Modular Monolith architecture, pairing a high-performance Express.js REST API "
        "with an autonomous BullMQ background worker and a React 18 single-page application (SPA)."
    )
    
    add_h2("Component Architecture Overview")
    comp_tbl = doc.add_table(rows=1, cols=3)
    comp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = comp_tbl.rows[0].cells
    hdr_cells[0].text = "Component"
    hdr_cells[1].text = "Technology Stack"
    hdr_cells[2].text = "Responsibility & Function"
    for c in hdr_cells:
        set_cell_shading(c, "1E293B")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    comp_data = [
        ("Frontend Web Portal", "React 18, Vite, TypeScript, Tailwind CSS, Lucide Icons, Recharts", "Modern responsive web dashboard, quick persona switching, real-time job status polling, interactive power controls, dynamic quota bars."),
        ("REST API Service", "Node.js 22, Express.js, TypeScript, Zod, Helmet, JWT, bcrypt", "Handles authentication, RBAC authorization, request CRUD, approval state machine transitions, quota validation, and job enqueuing."),
        ("Database Layer", "PostgreSQL 15, Prisma ORM", "ACID-compliant relational store managing users, roles, quotas, requests, provisioning jobs, active VMs, and immutable audit logs."),
        ("Queue & Cache", "Redis 7, BullMQ", "In-memory reliable message broker managing the 'provisioning' queue with retry policies, job state persistence, and worker coordination."),
        ("Provisioning Worker", "Dedicated Node.js Process, Child Process Sandbox, Terraform CLI", "Pulls jobs atomically from BullMQ, creates isolated execution sandboxes, runs Terraform stages (validate, plan, apply), captures logs, and updates DB."),
        ("Reverse Proxy", "Nginx 1.27 Alpine (in Docker)", "Serves compiled SPA assets on port 80/5173 and reverse-proxies /api/ requests to the API backend container.")
    ]
    for comp, tech, resp in comp_data:
        row_cells = comp_tbl.add_row().cells
        row_cells[0].text = comp
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].text = tech
        row_cells[2].text = resp
        for cell in row_cells:
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    add_h2("The 5-Stage Infrastructure Provisioning Lifecycle")
    doc.add_paragraph("When a developer submits a VM request, it transitions through a formally verified state machine:")
    
    lifecycle_steps = [
        ("Step 1: Submission & Quota Gate", "Developer submits request. Express API validates specs with Zod and checks database quotas. If approved quota would be exceeded, the request is rejected immediately with detailed excess metrics. Otherwise, status is set to PENDING_APPROVAL."),
        ("Step 2: Administrative Review", "An Enterprise Admin inspects pending requests in the Approvals queue. The Admin submits a decision with mandatory justification. Status moves to APPROVED (or REJECTED)."),
        ("Step 3: Background Job Enqueue", "Upon approval, the Status Transition Service atomically creates a ProvisioningJob record and enqueues a job payload to the BullMQ 'provisioning' queue in Redis."),
        ("Step 4: Terraform Sandbox Execution", "The background worker picks up the job, creates a unique working sandbox directory, and runs: (1) Validation -> (2) Planning -> (3) Applying. Execution output is continuously streamed into database logs."),
        ("Step 5: Inventory Registration & Notification", "Upon successful apply, a VirtualMachine record is created with simulated IP and MAC address, status transitions to ACTIVE, and an audit event is persisted.")
    ]
    for step, desc in lifecycle_steps:
        p = doc.add_paragraph(style='List Bullet')
        r1 = p.add_run(f"{step}: ")
        r1.font.bold = True
        p.add_run(desc)

    # 4. User Guide & Walkthrough
    add_h1("4. Step-by-Step User Guide (How to Use InfraPilot)")
    doc.add_paragraph(
        "InfraPilot features pre-configured demo personas accessible via one-click quick login buttons on the login screen:"
    )
    
    # Creds table
    cred_tbl = doc.add_table(rows=1, cols=4)
    cred_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdrs = cred_tbl.rows[0].cells
    c_hdrs[0].text = "Persona"
    c_hdrs[1].text = "Email"
    c_hdrs[2].text = "Password"
    c_hdrs[3].text = "Role & Access Scope"
    for c in c_hdrs:
        set_cell_shading(c, "2563EB")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    c_data = [
        ("Enterprise Admin", "admin@infrapilot.local", "AdminPass123!", "Global Fleet Visibility, Approvals, Quotas, Audit Logs, Job Retries"),
        ("Senior IT Developer", "developer@infrapilot.local", "DevPass123!", "Self-Service Requests, Personal Inventory, VM Power Actions"),
        ("DevOps Engineer", "developer2@infrapilot.local", "DevPass123!", "Personal Inventory (Isolated tenant - cannot see Developer 1's VMs)")
    ]
    for p_name, em, pw, scope in c_data:
        r_cells = cred_tbl.add_row().cells
        r_cells[0].text = p_name
        r_cells[0].paragraphs[0].runs[0].font.bold = True
        r_cells[1].text = em
        r_cells[2].text = pw
        r_cells[3].text = scope
        for cell in r_cells:
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    add_h2("Scenario A: Developer Workflow (Requesting a VM & Power Management)")
    dev_steps = [
        "1. Open http://localhost:5173/login in your web browser.",
        "2. Click the 'Developer Demo' quick button and click 'Sign In'.",
        "3. Notice the Developer Dashboard displays personal metrics (VM count, personal CPU usage).",
        "4. Click 'Request VM' in the sidebar.",
        "5. Enter VM Details: Name: 'dev-api-node', OS: 'Ubuntu 22.04 LTS', CPU: 2 Cores, RAM: 4 GB, Disk: 40 GB, Justification: 'API load testing'.",
        "6. Click 'Submit Request for Approval'. The request is created in PENDING_APPROVAL status.",
        "7. View VM Inventory: Click 'VM Inventory' in the sidebar to see currently active instances.",
        "8. Power Controls: Click 'Stop' or 'Start' on any instance to trigger simulated power operations with immediate audit logging."
    ]
    for s in dev_steps:
        doc.add_paragraph(s)

    add_h2("Scenario B: Administrator Workflow (Approval & Fleet Governance)")
    admin_steps = [
        "1. Log out and click the 'Admin Demo' quick button on the login screen.",
        "2. Observe the Admin Dashboard showing global metrics across all developers and infrastructure.",
        "3. Navigate to 'Approvals' in the sidebar: The pending request 'dev-api-node' appears at the top.",
        "4. Review specifications and click 'Approve'. Enter an approval comment and click 'Confirm Approval'.",
        "5. Navigate to 'Provisioning Jobs': Observe real-time progress as the BullMQ worker executes Terraform stages (VALIDATING -> PLANNING -> APPLYING -> COMPLETED).",
        "6. Click 'Logs' to inspect detailed Terraform execution output and generated JSON manifests.",
        "7. Navigate to 'Resource Quotas': View and adjust individual developer limits dynamically using interactive sliders.",
        "8. Navigate to 'Audit Logs': Inspect the immutable trail of all system operations with structured metadata."
    ]
    for s in admin_steps:
        doc.add_paragraph(s)

    # 5. How to Run the Project
    add_h1("5. How to Run and Deploy the Project")
    doc.add_paragraph("InfraPilot supports two execution modes: Docker Compose (recommended) or Local Development.")

    add_h2("Mode A: Docker Compose (All-in-One 5-Container Stack)")
    doc.add_paragraph("Run the complete stack with a single command from the project root:")
    p_code = doc.add_paragraph()
    r = p_code.add_run("docker compose up -d")
    r.font.name = 'Consolas'
    r.font.size = Pt(10)
    r.font.bold = True

    doc.add_paragraph(
        "Services started automatically:\n"
        "• Frontend Portal (Nginx): http://localhost:5173\n"
        "• REST API (Express): http://localhost:5000/api/v1\n"
        "• Provisioning Worker: BullMQ + Terraform CLI\n"
        "• PostgreSQL Database: localhost:5432\n"
        "• Redis Cache & Queue: localhost:6379"
    )

    add_h2("Mode B: Local Development")
    local_cmds = [
        "1. Database Sync: npm run db:push && npm run db:seed",
        "2. Start REST API: npm run dev:api (runs on port 5000)",
        "3. Start Background Worker: npm run dev:worker (in separate terminal)",
        "4. Start Frontend: npm run dev:frontend (runs on http://localhost:5173)",
        "5. Run Test Suite: npm test (executes 13 core business rule verification tests)"
    ]
    for c in local_cmds:
        doc.add_paragraph(c)

    # Save
    doc.save(DOCX_PATH)
    print(f"[DOCX] Word document generated successfully at: {DOCX_PATH}")

def generate_pdf():
    print("[PDF] Generating PDF (.pdf) documentation...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#0F172A'),
        alignment=1, # Center
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#475569'),
        alignment=1, # Center
        spaceAfter=15
    )

    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=19,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#2563EB'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'Bullet_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#1E293B'),
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=3
    )

    callout_style = ParagraphStyle(
        'Callout',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=4,
        spaceAfter=4
    )

    story = []

    # Title
    story.append(Paragraph("<b>ENTERPRISE ARCHITECTURE & USER MANUAL</b>", ParagraphStyle('TopSub', fontName='Helvetica-Bold', fontSize=9, textColor=colors.HexColor('#2563EB'), alignment=1, spaceAfter=4)))
    story.append(Paragraph("InfraPilot Management Platform", title_style))
    story.append(Paragraph("Self-Service Infrastructure Provisioning, Governance, & Lifecycle Automation", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563EB'), spaceAfter=12))

    # Section 1
    story.append(Paragraph("1. Executive Summary & Purpose of the Application", h1_style))
    story.append(Paragraph(
        "<b>InfraPilot</b> is a production-grade, self-service enterprise infrastructure portal designed to eliminate "
        "the classic tension between developer agility and IT operational governance. In traditional enterprise environments, "
        "software engineering teams often endure multi-day delays waiting for infrastructure teams to manually approve "
        "and provision virtual machines (VMs) through ticketing systems. Conversely, granting developers unfettered direct "
        "access to cloud provider consoles creates severe security risks, untracked configuration drift, and runaway costs.",
        body_style
    ))
    story.append(Paragraph(
        "InfraPilot bridges this gap by providing an intuitive, web-based self-service catalog where developers can "
        "request compute infrastructure with automated quota checks. Behind the scenes, requests flow through an enterprise "
        "approval workflow, trigger asynchronous background job queues, execute HashiCorp Terraform automation scripts in "
        "isolated sandboxes, and catalog provisioned virtual machines in an immutable inventory with real-time power controls.",
        body_style
    ))

    # Callout Table
    callout_data = [[Paragraph("<b>📌 CORE VALUE PROPOSITION:</b> InfraPilot empowers engineers to provision infrastructure on-demand while guaranteeing IT leaders 100% policy compliance, strict budget quota enforcement, zero double-approvals, and complete audit logging.", callout_style)]]
    callout_tbl = Table(callout_data, colWidths=[504])
    callout_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EFF6FF')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#93C5FD')),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
        ('RIGHTPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(callout_tbl)
    story.append(Spacer(1, 8))

    # Section 2
    story.append(Paragraph("2. Key Capabilities & Architectural Highlights", h1_style))
    features = [
        ("Self-Service Catalog", "Standardized VM templates (OS distributions, CPU, RAM, Disk, Region, Environment) validated via Zod schemas."),
        ("Resource Quota Enforcement", "Real-time server-side validation preventing developers from exceeding allocated CPU, RAM, disk, or VM count limits."),
        ("Strict Separation of Duties (RBAC)", "Distinct roles: DEVELOPER (request, track, manage personal VMs) and ADMIN (fleet oversight, quota tuning, audit logs, approvals). Developers cannot approve requests; admins cannot approve their own requests."),
        ("Asynchronous Job Processing", "Dedicated BullMQ engine powered by Redis. HTTP REST requests return immediately while long-running Terraform automation executes reliably in background workers."),
        ("Terraform IaC Sandbox", "Executes terraform init, validate, plan, and apply in isolated directories using the safe hashicorp/local provider to generate JSON infrastructure manifests without cloud charges."),
        ("Fleet Inventory & Audit Trail", "Real-time state tracking with interactive Start/Stop power controls and an append-only compliance audit log recording all logins, approvals, and configuration changes.")
    ]
    for name, desc in features:
        story.append(Paragraph(f"• <b>{name}</b>: {desc}", bullet_style))

    # Section 3
    story.append(Paragraph("3. System Architecture & How It Works", h1_style))
    story.append(Paragraph(
        "InfraPilot adopts a modern Decoupled Modular Monolith architecture, pairing a high-performance Express.js REST API "
        "with an autonomous BullMQ background worker and a React 18 single-page application (SPA).",
        body_style
    ))

    # Component Table
    comp_headers = [
        Paragraph("<b>Component</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=9, textColor=colors.white)),
        Paragraph("<b>Tech Stack</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=9, textColor=colors.white)),
        Paragraph("<b>Responsibility</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=9, textColor=colors.white))
    ]
    comp_rows = [comp_headers]
    comp_list = [
        ("Frontend Web Portal", "React 18, Vite, Tailwind CSS, Recharts", "Modern web dashboard, quick persona switching, real-time job status polling, interactive power controls, dynamic quota bars."),
        ("REST API Service", "Node.js 22, Express.js, TypeScript, Zod, JWT", "Handles authentication, RBAC authorization, request CRUD, approval state machine transitions, quota validation, and job enqueuing."),
        ("Database Layer", "PostgreSQL 15, Prisma ORM", "ACID-compliant relational store managing users, roles, quotas, requests, provisioning jobs, active VMs, and immutable audit logs."),
        ("Queue & Cache", "Redis 7, BullMQ", "In-memory reliable message broker managing the 'provisioning' queue with retry policies and worker coordination."),
        ("Provisioning Worker", "Dedicated Node.js, Terraform CLI", "Pulls jobs atomically from BullMQ, creates isolated execution sandboxes, runs Terraform stages (validate, plan, apply), captures logs, and updates DB."),
        ("Reverse Proxy", "Nginx 1.27 Alpine", "Serves compiled SPA assets on port 5173 and reverse-proxies /api/ requests to the API backend container.")
    ]
    for c, t, r in comp_list:
        comp_rows.append([
            Paragraph(f"<b>{c}</b>", body_style),
            Paragraph(t, body_style),
            Paragraph(r, body_style)
        ])

    comp_table = Table(comp_rows, colWidths=[110, 120, 274])
    comp_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0F172A')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(comp_table)
    story.append(Spacer(1, 8))

    # Section 4
    story.append(Paragraph("4. Step-by-Step User Guide (How to Use InfraPilot)", h1_style))
    story.append(Paragraph("InfraPilot includes built-in credentials with 1-click quick fill buttons on the login screen:", body_style))

    cred_headers = [
        Paragraph("<b>Persona</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=9, textColor=colors.white)),
        Paragraph("<b>Email</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=9, textColor=colors.white)),
        Paragraph("<b>Password</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=9, textColor=colors.white)),
        Paragraph("<b>Access Level</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=9, textColor=colors.white))
    ]
    cred_rows = [cred_headers]
    c_list = [
        ("Enterprise Admin", "admin@infrapilot.local", "AdminPass123!", "Global Fleet Visibility, Approvals, Quotas, Audit Logs, Job Retries"),
        ("Senior IT Developer", "developer@infrapilot.local", "DevPass123!", "Self-Service VM Requests, Personal Inventory, Power State"),
        ("DevOps Engineer", "developer2@infrapilot.local", "DevPass123!", "Personal Inventory (Tenant Isolation test)")
    ]
    for p, e, pw, sc in c_list:
        cred_rows.append([
            Paragraph(f"<b>{p}</b>", body_style),
            Paragraph(e, body_style),
            Paragraph(pw, body_style),
            Paragraph(sc, body_style)
        ])

    cred_table = Table(cred_rows, colWidths=[110, 140, 94, 160])
    cred_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2563EB')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(cred_table)
    story.append(Spacer(1, 8))

    story.append(Paragraph("Developer User Journey", h2_style))
    dev_walkthrough = [
        "1. Open http://localhost:5173/login and click 'Developer Demo' -> 'Sign In'.",
        "2. Note that personal metrics are displayed and admin-only menus are hidden.",
        "3. Click 'Request VM' in the sidebar.",
        "4. Quota Test: Select 16 Cores and 32GB RAM -> Red banner blocks submission exceeding maximum quota.",
        "5. Valid Request: Select 2 Cores, 4GB RAM, 40GB Disk, VM Name: 'dev-api-node', Justification: 'API load test' -> Status PENDING_APPROVAL.",
        "6. Power Actions: Navigate to 'VM Inventory' and click 'Stop' / 'Start' on active VMs to manage power state."
    ]
    for item in dev_walkthrough:
        story.append(Paragraph(item, bullet_style))

    story.append(Paragraph("Administrator User Journey", h2_style))
    admin_walkthrough = [
        "1. Log out and click 'Admin Demo' -> 'Sign In'.",
        "2. Fleet Dashboard shows global metrics across all developers and infrastructure.",
        "3. Click 'Approvals': Review 'dev-api-node' at the top of the queue and click 'Approve'.",
        "4. Click 'Provisioning Jobs': Observe real-time progress (VALIDATING -> PLANNING -> APPLYING -> COMPLETED).",
        "5. Inspect Logs: Click 'Logs' to see live Terraform execution output and generated JSON manifests.",
        "6. Adjust Quotas: Navigate to 'Resource Quotas' and adjust developer limits dynamically.",
        "7. Review Audit Logs: Inspect the immutable trail of all actions with structured event payloads."
    ]
    for item in admin_walkthrough:
        story.append(Paragraph(item, bullet_style))

    # Section 5
    story.append(Paragraph("5. How to Run the Application", h1_style))
    story.append(Paragraph(
        "<b>Single-Command Container Deployment (Docker Compose):</b><br/>"
        "Run <code>docker compose up -d</code> from the project root.<br/>"
        "• Web Portal: <b>http://localhost:5173</b><br/>"
        "• API Health Check: <b>http://localhost:5173/api/v1/health</b><br/>"
        "• Automated Test Suite: <code>npm test</code> (runs 13 Vitest verification tests)",
        body_style
    ))

    # Build PDF
    def add_header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#64748B'))
        # Header
        canvas.drawString(54, 750, "InfraPilot — Enterprise Self-Service Infrastructure Portal")
        canvas.drawRightString(612 - 54, 750, "Project Documentation")
        canvas.setStrokeColor(colors.HexColor('#CBD5E1'))
        canvas.setLineWidth(0.5)
        canvas.line(54, 744, 612 - 54, 744)
        # Footer
        canvas.line(54, 45, 612 - 54, 45)
        canvas.drawString(54, 32, "Confidential — For Technical Evaluation & Review")
        canvas.drawRightString(612 - 54, 32, f"Page {doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=add_header_footer, onLaterPages=add_header_footer)
    print(f"[PDF] PDF document generated successfully at: {PDF_PATH}")

if __name__ == '__main__':
    generate_docx()
    generate_pdf()
